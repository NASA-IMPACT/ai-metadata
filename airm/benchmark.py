"""Build the experiment query set from the external query benchmark.

``queries_benchmark.xlsx`` is a hand-curated benchmark: each row is a
natural-language question derived from a real published study, tagged with the
difficulty stratum, the single ground-truth CMR ``concept_id`` that answers it,
and the CMR metadata fields a reader needs to confirm the match. This module
turns that workbook into ``data/queries.yaml`` (the harness query set) and
guarantees every ground-truth collection is present in ``data/corpus.jsonl`` so
retrieval actually has the target to find.

Run::

    uv run python -m airm.benchmark            # default workbook at repo root
    uv run python -m airm.benchmark path.xlsx  # explicit workbook

It is a one-shot data-prep step, not part of the runtime harness — hence the
openpyxl dependency lives in dev deps.
"""

from __future__ import annotations

import re
from pathlib import Path

from . import QUERIES_PATH
from .cmr import extend_corpus
from .queries import Query, save_queries

DEFAULT_WORKBOOK = Path("queries_benchmark.xlsx")
_COLUMNS = ("source_file", "paper_title", "difficulty", "concept_id", "query", "fields")
# A real CMR collection concept-id: ``C`` + digits + ``-`` + provider token.
# Rows whose target couldn't be located carry placeholders (N/A, CMR_NOT_FOUND,
# UNKNOWN, blank) — those have no usable retrieval ground truth.
_CID_RE = re.compile(r"^C\d+-[A-Z0-9_]+$")


def _split_fields(raw: str | None) -> list[str]:
    """Comma-split a fields cell, ignoring commas nested in parentheses.

    The benchmark's harder rows annotate fields with parenthetical detail
    (e.g. ``Spatial resolution (~1.29 km, nadir)``); a naive ``split(",")``
    would shred those, so we only break at top-level commas.
    """
    if not raw:
        return []
    out: list[str] = []
    depth = 0
    buf: list[str] = []
    for ch in raw:
        if ch in "([":
            depth += 1
        elif ch in ")]":
            depth = max(0, depth - 1)
        if ch == "," and depth == 0:
            out.append("".join(buf).strip())
            buf = []
        else:
            buf.append(ch)
    if buf:
        out.append("".join(buf).strip())
    return [f for f in out if f]


def load_benchmark(workbook: Path = DEFAULT_WORKBOOK) -> tuple[list[Query], int]:
    """Parse the benchmark workbook into Query objects (order preserved).

    Returns ``(queries, n_dropped)`` where dropped rows are those with no
    resolvable CMR concept-id (placeholders) — they cannot anchor retrieval.
    """
    import openpyxl

    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb["Queries"]
    rows = ws.iter_rows(min_row=2, values_only=True)

    queries: list[Query] = []
    dropped = 0
    for i, row in enumerate(rows):
        rec = dict(zip(_COLUMNS, row))
        question = (rec.get("query") or "").strip()
        cid = (rec.get("concept_id") or "").strip()
        if not question:
            continue
        if not _CID_RE.match(cid):
            dropped += 1
            continue
        queries.append(
            Query(
                id=f"q{i:03d}",
                question=question,
                relevant=[cid],
                difficulty=(rec.get("difficulty") or "").strip(),
                fields=_split_fields(rec.get("fields")),
                paper=(rec.get("paper_title") or "").strip(),
            )
        )
    return queries, dropped


def build(workbook: Path = DEFAULT_WORKBOOK, out_path: Path = QUERIES_PATH) -> None:
    queries, dropped = load_benchmark(workbook)
    if dropped:
        # CLAUDE.md: never silently drop coverage — log it.
        print(f"dropped {dropped} query rows with no resolvable CMR concept-id")
    targets = sorted({cid for q in queries for cid in q.relevant})

    n_added, missing = extend_corpus(targets)
    print(f"corpus: +{n_added} ground-truth collections ({len(targets)} targets)")
    if missing:
        # CLAUDE.md: never silently drop coverage — surface unfetchable targets.
        # Their queries are unanswerable (target retired from CMR), so they would
        # only depress retrieval metrics with noise; drop them but report which.
        unresolved = set(missing)
        kept = [q for q in queries if not unresolved.intersection(q.relevant)]
        print(
            f"WARNING: {len(missing)} target(s) not in CMR: {sorted(unresolved)} — "
            f"dropping {len(queries) - len(kept)} query(ies) pointing at them"
        )
        queries = kept

    save_queries(queries, out_path)
    by_diff: dict[str, int] = {}
    for q in queries:
        by_diff[q.difficulty] = by_diff.get(q.difficulty, 0) + 1
    print(f"queries: wrote {len(queries)} to {out_path} ({by_diff})")


if __name__ == "__main__":  # pragma: no cover - manual data-prep entrypoint
    import sys

    wb = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    build(wb)
