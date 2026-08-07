"""The evaluation query set and its ground truth.

Two halves:

* **SME queries** — expert-authored, read from ``sample_sme_queries.xlsx``. Kept
  verbatim and treated as the held-out expert slice.
* **Synthetic queries** — generated from indexed records in Step 9, so their
  ground truth is in-corpus by construction.

This module owns the SME half. Everything it drops is written to
``data/gt_dropped.json`` with a reason; nothing disappears silently.
"""

from __future__ import annotations

import json
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterable

from .config import GT_DROPPED_PATH, GT_PATH, SME_XLSX_PATH

#: A CMR collection concept-id: ``C`` + digits + ``-`` + provider.
CONCEPT_ID_RE = re.compile(r"^C\d+-[A-Z0-9_]+$")

#: Unicode dashes the spreadsheet uses in place of ASCII hyphen-minus. Three
#: cells contain U+2011 NON-BREAKING HYPHEN, which makes the id look correct to a
#: human and fail every string comparison.
DASHES = "‐‑‒–—−"


@dataclass
class Query:
    """One evaluation query and the collections that should answer it."""

    query_id: str
    text: str
    expected_concept_ids: list[str]
    source: str  # "sme" | "synthetic"
    topic: str | None = None
    expected_titles: list[str] = field(default_factory=list)

    def to_json(self) -> str:
        return json.dumps(asdict(self), ensure_ascii=False)


# --------------------------------------------------------------------------- #
# Spreadsheet parsing.
# --------------------------------------------------------------------------- #


def normalise_concept_id(token: str) -> str:
    """Strip whitespace and fold Unicode dashes down to ASCII hyphen-minus."""
    cleaned = unicodedata.normalize("NFKC", token).strip()
    for dash in DASHES:
        cleaned = cleaned.replace(dash, "-")
    return "".join(cleaned.split())


def split_concept_ids(cell: str) -> list[tuple[str, str]]:
    """Split a concept-id cell into ``(original text, normalised token)`` pairs."""
    if not cell:
        return []
    out = []
    for part in cell.split(","):
        raw = part.strip()
        token = normalise_concept_id(part)
        if token:
            out.append((raw, token))
    return out


def read_sme_rows(path: Path = SME_XLSX_PATH) -> list[dict]:
    """Read ``(query, expected_dataset, concept_id_cell)`` rows from the sheet."""
    from openpyxl import load_workbook

    if not path.exists():
        raise FileNotFoundError(f"SME query workbook not found at {path}")

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        raise ValueError(f"{path} is empty")

    header = [str(c or "").strip().lower() for c in rows[0]]
    try:
        i_query = header.index("query")
        i_dataset = header.index("expected dataset")
        i_ids = next(i for i, h in enumerate(header) if "concept id" in h)
    except (ValueError, StopIteration) as exc:
        raise ValueError(f"unexpected header in {path}: {rows[0]}") from exc

    out = []
    for row in rows[1:]:
        text = str(row[i_query] or "").strip()
        if not text:
            continue
        out.append(
            {
                "text": text,
                "dataset": str(row[i_dataset] or "").strip(),
                "ids_cell": str(row[i_ids] or "").strip(),
            }
        )
    return out


# --------------------------------------------------------------------------- #
# Ground truth.
# --------------------------------------------------------------------------- #


def build_sme_queries(
    path: Path = SME_XLSX_PATH,
    *,
    resolve: bool = True,
) -> tuple[list[Query], dict]:
    """Build the SME query set and a report of everything dropped.

    Rows sharing a query string are merged into one query whose expected set is
    the union of their concept-ids -- the sheet splits a single research question
    across several rows, one per dataset family, and scoring them as separate
    queries would count the same question up to four times.

    Two classes of token are removed, and both are recorded:

    ``not_a_concept_id``
        Cells that are not concept-ids at all (a granule short-name, the literal
        note ``Not in CMR``).
    ``retired``
        Well-formed ids CMR no longer resolves. Per the study's decision these
        are dropped rather than remapped to a successor collection, so ground
        truth only ever contains collections CMR returns today.

    Set ``resolve=False`` to skip the network check (offline tests).
    """
    rows = read_sme_rows(path)

    merged: dict[str, dict] = {}
    for row in rows:
        entry = merged.setdefault(
            row["text"], {"ids": [], "titles": [], "raw": {}, "order": len(merged)}
        )
        for raw, token in split_concept_ids(row["ids_cell"]):
            if token not in entry["ids"]:
                entry["ids"].append(token)
                # Keep the cell text so the dropped-token report shows what a
                # reader would actually find in the spreadsheet, not the
                # whitespace-stripped form ("Not in CMR", not "NotinCMR").
                entry["raw"][token] = raw
        if row["dataset"] and row["dataset"] not in entry["titles"]:
            entry["titles"].append(row["dataset"])

    malformed: dict[str, list[str]] = {}
    for text, entry in merged.items():
        bad = [entry["raw"][t] for t in entry["ids"] if not CONCEPT_ID_RE.match(t)]
        if bad:
            malformed[text] = bad
        entry["ids"] = [t for t in entry["ids"] if CONCEPT_ID_RE.match(t)]

    all_ids = sorted({cid for e in merged.values() for cid in e["ids"]})

    retired: list[str] = []
    if resolve:
        from .cmr import missing_concept_ids

        retired = missing_concept_ids(all_ids)

    retired_set = set(retired)
    queries: list[Query] = []
    emptied: list[str] = []
    for text, entry in sorted(merged.items(), key=lambda kv: kv[1]["order"]):
        kept = [cid for cid in entry["ids"] if cid not in retired_set]
        if not kept:
            emptied.append(text)
            continue
        queries.append(
            Query(
                query_id=f"sme-{len(queries) + 1:03d}",
                text=text,
                expected_concept_ids=kept,
                source="sme",
                expected_titles=entry["titles"],
            )
        )

    kept_unique = ground_truth_ids(queries)
    malformed_unique = sorted({t for tokens in malformed.values() for t in tokens})

    report = {
        "source_file": str(path),
        "rows_read": len(rows),
        "distinct_queries": len(merged),
        "queries_kept": len(queries),
        # Unique counts. A raw sum over queries would double-count the ids that
        # legitimately answer more than one research question -- three do.
        "concept_id_tokens_seen": len(all_ids) + len(malformed_unique),
        "concept_ids_well_formed": len(all_ids),
        "concept_ids_kept": len(kept_unique),
        "concept_id_slots_across_queries": sum(len(q.expected_concept_ids) for q in queries),
        "kept_concept_ids": kept_unique,
        "dropped": {
            "not_a_concept_id": malformed,
            "retired": sorted(retired),
        },
        "queries_dropped_for_having_no_resolvable_ids": emptied,
    }
    return queries, report


# --------------------------------------------------------------------------- #
# Persistence.
# --------------------------------------------------------------------------- #


def write_queries(queries: Iterable[Query], path: Path = GT_PATH) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w") as fh:
        for q in queries:
            fh.write(q.to_json() + "\n")
    tmp.replace(path)
    return path


def load_queries(path: Path = GT_PATH) -> list[Query]:
    if not path.exists():
        raise FileNotFoundError(
            f"No query set at {path}. Run `uv run python -m airm.queries --build` first."
        )
    with path.open() as fh:
        return [Query(**json.loads(line)) for line in fh if line.strip()]


def load_eval_queries() -> list[Query]:
    """The full evaluation set: SME + synthetic when built, SME alone otherwise.

    ``airm.synth`` writes the combined set to ``QUERIES_PATH``; before it has
    run, only the SME ground truth at ``GT_PATH`` exists. Experiment 2 must go
    through this loader -- loading ``GT_PATH`` directly would silently run the
    11 expert queries and skip the 120 synthetic ones, and every gate would
    still pass.
    """
    from .config import QUERIES_PATH

    if QUERIES_PATH.exists():
        return load_queries(QUERIES_PATH)
    return load_queries()


def write_report(report: dict, path: Path = GT_DROPPED_PATH) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n")
    return path


def ground_truth_ids(queries: Iterable[Query]) -> list[str]:
    """Every concept-id the query set expects, deduplicated and sorted."""
    return sorted({cid for q in queries for cid in q.expected_concept_ids})


def build(path: Path = SME_XLSX_PATH) -> tuple[list[Query], dict]:
    """Build, persist and summarise the SME half of the query set."""
    queries, report = build_sme_queries(path)
    write_queries(queries)
    write_report(report)
    return queries, report


# --------------------------------------------------------------------------- #
# Synthetic queries.
# --------------------------------------------------------------------------- #

GENERATION_SYSTEM = """\
You write realistic Earth-science data-discovery questions for evaluating \
metadata search systems.

You are shown one NASA CMR collection. Write a question a domain scientist would \
plausibly ask that this collection helps answer.

Rules:
- Ask a research question, not a lookup. "How did sea-ice thickness in the Kara \
Sea change each March from 2020 to 2025?" is good; "Find the CryoSat-2 sea ice \
product" is not.
- NEVER name the dataset, its short name, its DOI, its concept id, its version, \
or its data centre. The question must be answerable only by understanding what \
the data measures.
- Ground it in the science: name the phenomenon, a plausible region and a \
plausible time period.
- One or two sentences.

Return JSON: {"query": "<the question>"}"""


def _generation_prompt(payload: dict) -> str:
    lines = [f"Title: {payload.get('title', '')}"]
    if payload.get("topic"):
        lines.append(f"Science topic: {payload['topic']}")
    if payload.get("science_keywords"):
        lines.append("Keywords: " + "; ".join(payload["science_keywords"][:6]))
    if payload.get("variables"):
        lines.append("Measures: " + ", ".join(payload["variables"][:8]))
    if payload.get("platforms"):
        lines.append(
            "Instruments: "
            + "; ".join(
                f"{p['platform']} ({', '.join(p['instruments'])})"
                if p.get("instruments")
                else p["platform"]
                for p in payload["platforms"][:4]
            )
        )
    if payload.get("bbox"):
        b = payload["bbox"]
        lines.append(f"Spatial: {b['west']} to {b['east']} lon, {b['south']} to {b['north']} lat")
    if payload.get("temporal"):
        t = payload["temporal"]
        lines.append(f"Temporal: {t['begin']} to {t.get('end', 'unspecified')}")
    if payload.get("summary"):
        lines.append(f"Abstract: {payload['summary'][:900]}")
    return "\n".join(lines)


#: A verbatim title run this long is quoting, not describing.
TITLE_NGRAM = 5

_WORD_RE = re.compile(r"[a-z0-9]+")


def _words(text: str) -> list[str]:
    return _WORD_RE.findall(text.lower())


def _leaked_identifiers(text: str, payload: dict) -> list[str]:
    """Identifiers a generated query must not quote.

    A query naming the short name, DOI or concept id turns retrieval into string
    matching: every format contains that literal, so every format scores alike
    and the effect under measurement disappears.

    The title check compares *word sequences* on both sides rather than matching
    a filtered phrase against the raw text. Filtering short words out of the
    title only ("ATLAS ICESat Land Vegetation Height") and then searching for it
    in a query that kept them ("ATLAS ICESat Land and Vegetation Height") never
    matches, so a title quoted with its stopwords -- the likeliest way a model
    quotes one -- slipped straight through.
    """
    lowered = text.lower()
    leaks = []
    for key in ("concept_id", "short_name", "doi"):
        value = payload.get(key)
        if value and str(value).lower() in lowered:
            leaks.append(f"{key}={value}")

    title_words = _words(payload.get("title", ""))
    query_words = _words(text)
    if len(title_words) >= TITLE_NGRAM and len(query_words) >= TITLE_NGRAM:
        query_grams = {
            tuple(query_words[i : i + TITLE_NGRAM])
            for i in range(len(query_words) - TITLE_NGRAM + 1)
        }
        for i in range(len(title_words) - TITLE_NGRAM + 1):
            gram = tuple(title_words[i : i + TITLE_NGRAM])
            if gram in query_grams:
                leaks.append(f"title-phrase={' '.join(gram)!r}")
                break
    return leaks


def parse_generated(text: str) -> str | None:
    """Pull the query string out of a model response, tolerating stray prose."""
    text = text.strip()
    if text.startswith("```"):
        text = text.split("```")[1].lstrip("json").strip()
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start == -1 or end <= start:
            return None
        try:
            data = json.loads(text[start : end + 1])
        except json.JSONDecodeError:
            return None
    value = data.get("query") if isinstance(data, dict) else None
    return value.strip() if isinstance(value, str) and value.strip() else None


if __name__ == "__main__":  # pragma: no cover - CLI
    import argparse

    parser = argparse.ArgumentParser(description="Build the SME query ground truth.")
    parser.add_argument("--build", action="store_true")
    args = parser.parse_args()

    qs, rep = build() if args.build else build_sme_queries()

    print(f"rows read                : {rep['rows_read']}")
    print(f"distinct queries         : {rep['distinct_queries']}")
    print(f"queries kept             : {rep['queries_kept']}")
    print(f"concept-id tokens seen   : {rep['concept_id_tokens_seen']} unique")
    print(f"  well-formed            : {rep['concept_ids_well_formed']}")
    print(f"  kept (resolve in CMR)  : {rep['concept_ids_kept']} unique"
          f"  ({rep['concept_id_slots_across_queries']} slots across queries)")
    print(f"dropped (not an id)      : {sorted({t for v in rep['dropped']['not_a_concept_id'].values() for t in v})}")
    print(f"dropped (retired in CMR) : {rep['dropped']['retired']}")
    if rep["queries_dropped_for_having_no_resolvable_ids"]:
        print("queries with no resolvable ids:")
        for text in rep["queries_dropped_for_having_no_resolvable_ids"]:
            print(f"  - {text[:90]}")
    if args.build:
        print(f"\nwrote {GT_PATH}\nwrote {GT_DROPPED_PATH}")
    for q in qs:
        print(f"  {q.query_id}  {len(q.expected_concept_ids):>2} ids  {q.text[:78]}")
